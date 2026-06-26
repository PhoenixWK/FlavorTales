"""Generate stress-test-strategy.docx in the same folder as this script."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
import os

doc = Document()

# ── Page margins ────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(2)

# ── Helpers ─────────────────────────────────────────────────────────────────
def heading(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79) if level == 1 else RGBColor(0x2E, 0x74, 0xB5)
    return p

def para(text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text).font.size = Pt(11)
    return p

def code_block(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
    return p

def table_row(tbl, cells, bold=False, bg=None):
    row = tbl.add_row()
    for i, text in enumerate(cells):
        cell = row.cells[i]
        cell.text = text
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = bold
            run.font.size = Pt(10)
    return row

# ── Title ────────────────────────────────────────────────────────────────────
title = doc.add_heading("Chiến Lược Stress Test – FlavorTales Tourist", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run("Target: https://flavortales.site  |  Ngày: 02/04/2026").font.color.rgb = RGBColor(0x59, 0x59, 0x59)

doc.add_paragraph()

# ── 1. Mục tiêu ──────────────────────────────────────────────────────────────
heading("1. Mục Tiêu")
para(
    "Xác định ngưỡng chịu tải tối đa (breaking point) của server FlavorTales "
    "dưới traffic tourist ẩn danh (anonymous). Cụ thể:"
)
bullet("Tìm số lượng concurrent user tại đó error rate vượt 5%.")
bullet("Ghi nhận P95 response time tại từng mốc tải.")
bullet("Quan sát thứ tự tài nguyên nào bị bão hoà trước (CPU, DB, network).")

# ── 2. Phạm vi ───────────────────────────────────────────────────────────────
heading("2. Phạm Vi Test")
para(
    "Chỉ test các endpoint dành cho tourist ẩn danh — không cần JWT, không cần đăng nhập:"
)

tbl = doc.add_table(rows=1, cols=3)
tbl.style = "Table Grid"
hdr = tbl.rows[0].cells
for i, t in enumerate(["Method", "Endpoint", "Mô tả"]):
    hdr[i].text = t
    for run in hdr[i].paragraphs[0].runs:
        run.bold = True
        run.font.size = Pt(10)

rows_data = [
    ("POST",   "/api/tourist/sessions",              "Tạo phiên ẩn danh mới"),
    ("GET",    "/api/tourist/sessions/{sessionId}",  "Kiểm tra phiên còn hiệu lực"),
    ("GET",    "/api/poi",                           "Tải danh sách POI (bản đồ)"),
    ("GET",    "/api/audio/poi/{poiId}",             "Tải danh sách audio của POI"),
    ("POST",   "/api/poi/{poiId}/like",              "Like một POI (kèm X-Session-Id)"),
    ("DELETE", "/api/tourist/sessions/{sessionId}",  "Kết thúc phiên"),
]
for r in rows_data:
    table_row(tbl, r)

doc.add_paragraph()

# ── 3. Công cụ ───────────────────────────────────────────────────────────────
heading("3. Công Cụ – Locust")
para("Locust là công cụ load testing viết bằng Python, có giao diện Web UI real-time.")
bullet("Cài đặt: python -m pip install locust")
bullet("Chạy: python -m locust -f test/locustfile.py --host=https://flavortales.site")
bullet("Mở Web UI tại: http://localhost:8089")
para(
    "Locust cho phép tăng/giảm số lượng virtual user ngay trong lúc test đang chạy "
    "thông qua Web UI, không cần dừng và chạy lại.",
    bold=False,
)

# ── 4. Mô hình virtual user ───────────────────────────────────────────────────
heading("4. Mô Hình Virtual User")
para(
    "Mỗi virtual user (VU) mô phỏng một tourist ẩn danh với luồng hành vi sau:"
)
bullet("[on_start] Tạo session → Tải danh sách POI → lưu sessionId và danh sách poiId")
bullet("[task × 5]  GET /api/poi              — duyệt/pan bản đồ (tần suất cao nhất)")
bullet("[task × 3]  GET /api/audio/poi/{id}  — click POI, nghe audio narration")
bullet("[task × 2]  POST /api/poi/{id}/like  — like một POI ngẫu nhiên")
bullet("[task × 1]  GET /api/tourist/sessions/{id} — kiểm tra session còn hiệu lực")
bullet("[on_stop]   DELETE /api/tourist/sessions/{id} — kết thúc phiên")
doc.add_paragraph()
para("Thời gian nghỉ giữa các request: 1 – 3 giây (realistic browser behaviour).")

# ── 5. Chiến lược ramp-up ────────────────────────────────────────────────────
heading("5. Chiến Lược Ramp-Up")
para(
    "Tăng dần số lượng VU, quan sát 2–3 phút mỗi mốc trước khi tăng tiếp. "
    "Tốc độ spawn: 5 users/giây."
)

tbl2 = doc.add_table(rows=1, cols=4)
tbl2.style = "Table Grid"
hdr2 = tbl2.rows[0].cells
for i, t in enumerate(["Giai đoạn", "Số VU", "Thời gian quan sát", "Mục tiêu"]):
    hdr2[i].text = t
    for run in hdr2[i].paragraphs[0].runs:
        run.bold = True
        run.font.size = Pt(10)

ramp_data = [
    ("Warm-up",      "10",   "2 phút", "Xác nhận tất cả endpoint trả 200 OK"),
    ("Nhẹ",          "50",   "3 phút", "Baseline: RPS và latency ổn định"),
    ("Trung bình",   "100",  "3 phút", "Quan sát response time P95"),
    ("Cao",          "200",  "3 phút", "Bắt đầu xuất hiện dấu hiệu căng thẳng?"),
    ("Rất cao",      "500",  "3 phút", "Tìm ngưỡng error rate > 5%"),
    ("Cực đại",      "1000+","đến sập","Breaking point — ghi lại số VU chính xác"),
]
for r in ramp_data:
    table_row(tbl2, r)

doc.add_paragraph()

# ── 6. Tiêu chí breaking point ───────────────────────────────────────────────
heading("6. Tiêu Chí Xác Định Breaking Point")
para("Server được coi là đã đạt breaking point khi ít nhất một trong các điều kiện sau xảy ra:")
bullet("Error rate vượt 5% (HTTP 4xx do server overload, 502, 503, 504, hoặc timeout).")
bullet("P95 response time vượt 5.000 ms (5 giây).")
bullet("RPS (requests/second) ngừng tăng dù vẫn thêm VU — dấu hiệu server đã bão hoà.")
bullet("Connection refused hoặc TLS handshake timeout liên tục.")

# ── 7. Cách đọc Locust Web UI ────────────────────────────────────────────────
heading("7. Cách Đọc Kết Quả Trên Locust Web UI")

tbl3 = doc.add_table(rows=1, cols=2)
tbl3.style = "Table Grid"
hdr3 = tbl3.rows[0].cells
for i, t in enumerate(["Chỉ số", "Ý nghĩa"]):
    hdr3[i].text = t
    for run in hdr3[i].paragraphs[0].runs:
        run.bold = True
        run.font.size = Pt(10)

metrics = [
    ("RPS (Requests/s)",        "Số request server xử lý được mỗi giây — tăng plateau = bão hoà"),
    ("Failures/s",              "Số request lỗi mỗi giây — tăng đột biến = server đang sập"),
    ("Median (50th %ile)",      "50% request hoàn thành trong thời gian này"),
    ("95th percentile (P95)",   "95% request hoàn thành trong thời gian này — chỉ số quan trọng nhất"),
    ("99th percentile (P99)",   "Tail latency — trải nghiệm của 1% user chậm nhất"),
    ("Users",                   "Số VU đang chạy hiện tại"),
    ("Charts tab",              "Biểu đồ real-time: RPS + Response time + Users cùng trục thời gian"),
]
for r in metrics:
    table_row(tbl3, r)

doc.add_paragraph()

# ── 8. Lưu ý an toàn ────────────────────────────────────────────────────────
heading("8. Lưu Ý An Toàn")
bullet("Chạy test vào giờ thấp điểm (ban đêm, sáng sớm) để tránh ảnh hưởng user thật.")
bullet("Bắt đầu từ 10 VU, không nhảy thẳng lên 500.")
bullet("Luôn kiểm tra Spring Boot logs phía server song song với Locust UI.")
bullet("Nếu server sập, đợi khởi động lại trước khi chạy tiếp.")
bullet("Không test các endpoint có rate-limit (forgot-password, register) trong bài này.")

# ── 9. Câu lệnh thực thi ─────────────────────────────────────────────────────
heading("9. Câu Lệnh Thực Thi")
para("Chạy lệnh sau từ thư mục gốc của workspace:")
code_block("python -m locust -f test/locustfile.py --host=https://flavortales.site")
para("Sau đó mở trình duyệt tại:")
code_block("http://localhost:8089")
para("Điền vào form:")
bullet("Number of users: 10 (warm-up)")
bullet("Spawn rate: 5")
bullet("Host: https://flavortales.site (đã được điền sẵn)")
para("Nhấn Start Swarming và quan sát. Tăng dần số VU theo bảng tại mục 5.")

# ── Save ─────────────────────────────────────────────────────────────────────
out_path = os.path.join(os.path.dirname(__file__), "stress-test-strategy.docx")
doc.save(out_path)
print(f"Saved: {out_path}")
