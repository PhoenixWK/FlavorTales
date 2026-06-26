"""Generate tourist-04-03-2026.docx evaluation report."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()

# ── Margins ─────────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(2)

# ── Colour palette ───────────────────────────────────────────────────────────
DARK_BLUE  = RGBColor(0x1F, 0x4E, 0x79)
MID_BLUE   = RGBColor(0x2E, 0x74, 0xB5)
GREEN      = RGBColor(0x37, 0x86, 0x1E)
RED        = RGBColor(0xC0, 0x00, 0x00)
ORANGE     = RGBColor(0xC5, 0x5A, 0x11)
GREY_TEXT  = RGBColor(0x59, 0x59, 0x59)

# ── Helpers ──────────────────────────────────────────────────────────────────
def heading(text, level=1):
    p = doc.add_heading(text, level=level)
    colour = DARK_BLUE if level == 1 else MID_BLUE
    for run in p.runs:
        run.font.color.rgb = colour
    return p

def para(text, bold=False, size=11, colour=None, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Inches(0.3)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if colour:
        run.font.color.rgb = colour
    return p

def bullet(text, colour=None):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.size = Pt(11)
    if colour:
        run.font.color.rgb = colour
    return p

def shade_cell(cell, hex_fill):
    """Apply background colour to a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)

def add_table(headers, rows_data, widths=None, header_fill='2E74B5'):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = 'Table Grid'
    if widths:
        for i, w in enumerate(widths):
            tbl.columns[i].width = Cm(w)
    hdr = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        shade_cell(hdr[i], header_fill)
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row_data in rows_data:
        row = tbl.add_row()
        for i, (text, fill, txt_colour) in enumerate(row_data):
            cell = row.cells[i]
            cell.text = str(text)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if fill:
                shade_cell(cell, fill)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(10)
                if txt_colour:
                    run.font.color.rgb = txt_colour
    return tbl

# ── NFR thresholds (from system_design_nfr.docx §1.1) ───────────────────────
NFR_API_MS        = 5000   # API Response Time < 5000 ms
NFR_CONCURRENT    = 100    # Concurrent tourist users (minimum)
NFR_ERROR_RATE    = 5.0    # Breaking point threshold (our test criterion)

# ── Test data (extracted from Locust HTML reports) ───────────────────────────
results = [
    {
        'stage': 'Warm-up',    'users': 10,  'total_req': 177,
        'avg_ms': 153,  'p95_ms': 270, 'p99_ms': 300,
        'rps': 1.43,   'num_failures': 0,  'error_rate_pct': 0.0,
    },
    {
        'stage': 'Nhẹ',        'users': 50,  'total_req': 1136,
        'avg_ms': 104,  'p95_ms': 160, 'p99_ms': 220,
        'rps': 6.23,   'num_failures': 0,  'error_rate_pct': 0.0,
    },
    {
        'stage': 'Trung bình', 'users': 100, 'total_req': 2257,
        'avg_ms': 99,   'p95_ms': 150, 'p99_ms': 240,
        'rps': 12.36,  'num_failures': 0,  'error_rate_pct': 0.0,
    },
    {
        'stage': 'Cao',        'users': 200, 'total_req': 4272,
        'avg_ms': 90,   'p95_ms': 110, 'p99_ms': 130,
        'rps': 23.45,  'num_failures': 0,  'error_rate_pct': 0.0,
    },
    {
        'stage': 'Rất cao',    'users': 500, 'total_req': 8767,
        'avg_ms': 91,   'p95_ms': 110, 'p99_ms': 140,
        'rps': 48.36,  'num_failures': 0,  'error_rate_pct': 0.0,
    },
    {
        'stage': 'Cực đại (1)',  'users': 905, 'total_req': 10651,
        'avg_ms': 129,  'p95_ms': 200, 'p99_ms': 810,
        'rps': 58.59,  'num_failures': 1,  'error_rate_pct': 0.01,
    },
    {
        'stage': 'Cực đại (2)',  'users': 895, 'total_req': 10843,
        'avg_ms': 118,  'p95_ms': 190, 'p99_ms': 260,
        'rps': 59.73,  'num_failures': 0,  'error_rate_pct': 0.0,
    },
]

def pass_fail_colour(ok):
    return ('E2EFDA', GREEN) if ok else ('FFE7E7', RED)

# ════════════════════════════════════════════════════════════════════════════ #
#  DOCUMENT                                                                   #
# ════════════════════════════════════════════════════════════════════════════ #

# ── Title ────────────────────────────────────────────────────────────────────
title = doc.add_heading('BÁO CÁO ĐÁNH GIÁ KẾT QUẢ STRESS TEST – TOURIST', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.color.rgb = DARK_BLUE

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run('FlavorTales  |  Target: https://flavortales.site  |  Ngày test: 03-04-2026')
run.font.color.rgb = GREY_TEXT
run.font.size = Pt(11)

doc.add_paragraph()

# ── 1. Tổng quan ─────────────────────────────────────────────────────────────
heading('1. Tổng Quan')
para(
    'Báo cáo này đánh giá kết quả 7 phiên stress test được thực hiện ngày 03-04-2026 '
    'trên môi trường production của hệ thống FlavorTales tại địa chỉ https://flavortales.site. '
    'Đối tượng test là tourist ẩn danh (anonymous) — không yêu cầu xác thực JWT. '
    'Mục tiêu: xác định ngưỡng chịu tải tối đa và đánh giá mức độ đáp ứng Yêu Cầu Phi Chức Năng (NFR) '
    'trong tài liệu system_design_nfr.docx.'
)

# ── 2. Yêu cầu NFR áp dụng ───────────────────────────────────────────────────
heading('2. Yêu Cầu NFR Áp Dụng (§1.1 – Performance Requirements)')
para('Các ngưỡng hiệu năng được trích từ tài liệu System Design NFR §1.1:', bold=True)
doc.add_paragraph()

nfr_rows = [
    [('API Response Time (P95/P99)', None, None), ('< 5.000 ms', None, None),  ('Điều kiện network bình thường', None, None)],
    [('Concurrent Tourist Users',    None, None), ('≥ 100 người đồng thời',     None, None), ('Truy cập bản đồ & phát audio',  None, None)],
    [('Error Rate (breaking point)', None, None), ('> 5% = breaking point',     None, None), ('Ngưỡng xác định server sập',    None, None)],
    [('POI Data Loading',            None, None), ('≤ 4.000 ms',                None, None), ('Sau khi tap vào marker',        None, None)],
    [('Audio Playback start',        None, None), ('≤ 4.000 ms',                None, None), ('Sau khi nhấn Play',             None, None)],
]
add_table(['Tiêu chí', 'Ngưỡng NFR', 'Điều kiện'], nfr_rows)
doc.add_paragraph()

# ── 3. Phương pháp & cấu hình test ──────────────────────────────────────────
heading('3. Phương Pháp & Cấu Hình Test')
add_table(
    ['Tham số', 'Giá trị'],
    [
        [('Công cụ', None, None),         ('Locust 2.43.4', None, None)],
        [('Giao thức', None, None),       ('HTTPS (TLS 1.3)', None, None)],
        [('Wait time/user', None, None),  ('1 – 3 giây (random)', None, None)],
        [('Spawn rate', None, None),      ('5 users/giây', None, None)],
        [('Luồng tourist', None, None),   ('POST session → GET /api/poi → GET /api/audio/poi/{id} → POST like → DELETE session', None, None)],
        [('Task weights', None, None),    ('GET /api/poi ×5 | GET audio ×3 | POST like ×2 | GET session ×1', None, None)],
    ],
    header_fill='1F4E79'
)
doc.add_paragraph()

# ── 4. Kết quả tổng hợp ──────────────────────────────────────────────────────
heading('4. Kết Quả Tổng Hợp Theo Giai Đoạn')

hdr4 = ['Giai đoạn', 'Users', 'Tổng req', 'Avg (ms)', 'P95 (ms)', 'P99 (ms)', 'RPS', 'Lỗi', 'Error %', 'NFR P95']
rows4 = []
for r in results:
    p95_ok = r['p95_ms'] < NFR_API_MS
    fill_p95, col_p95 = pass_fail_colour(p95_ok)
    err_ok = r['error_rate_pct'] < NFR_ERROR_RATE
    fill_err, col_err = pass_fail_colour(err_ok)
    rows4.append([
        (r['stage'],                                    None,     None),
        (str(r['users']),                               None,     None),
        (f"{r['total_req']:,}",                         None,     None),
        (f"{r['avg_ms']}",                              None,     None),
        (f"{r['p95_ms']}",                              fill_p95, col_p95),
        (f"{r['p99_ms']}",                              None,     None),
        (f"{r['rps']:.2f}",                             None,     None),
        (str(r['num_failures']),                        None,     None),
        (f"{r['error_rate_pct']:.2f}%",                 fill_err, col_err),
        ('✔ ĐẠT' if p95_ok else '✘ KHÔNG ĐẠT',         fill_p95, col_p95),
    ])

add_table(hdr4, rows4, header_fill='2E74B5')
doc.add_paragraph()
para('Chú thích: ô xanh lá = đạt ngưỡng NFR, ô đỏ = vượt ngưỡng.', size=9, colour=GREY_TEXT)
doc.add_paragraph()

# ── 5. Đánh giá từng tiêu chí NFR ──────────────────────────────────────────
heading('5. Đánh Giá Từng Tiêu Chí NFR')

# 5.1 API Response Time
heading('5.1. API Response Time (NFR: < 5.000 ms)', level=2)
para(
    'Tất cả 7 giai đoạn test đều có P95 response time dưới 300 ms, '
    'thấp hơn rất nhiều so với ngưỡng NFR 5.000 ms. '
    'Ngay cả ở tải cực đại (~900 concurrent users), P95 chỉ đạt 200 ms và P99 đạt 810 ms — '
    'vẫn nằm trong giới hạn cho phép.'
)
bullet('Giai đoạn tốt nhất: Cao (200 users) — P95 = 110 ms')
bullet('Giai đoạn cao nhất: Warm-up — P95 = 270 ms (do cold-start JVM/Redis)')
bullet('Kết luận: ✔ ĐẠT — Vượt trội so với yêu cầu NFR', colour=GREEN)
doc.add_paragraph()

# 5.2 Concurrent Users
heading('5.2. Concurrent Tourist Users (NFR: ≥ 100 đồng thời)', level=2)
para(
    'Hệ thống được test thành công ở mức 100 users (giai đoạn Trung bình) với error rate 0%, '
    'P95 = 150 ms. Hơn nữa, server duy trì hoạt động ổn định lên đến ~900 concurrent users '
    'mà không có sự cố nghiêm trọng.'
)
bullet('Tại 100 users: Error rate = 0%, P95 = 150 ms — đáp ứng hoàn toàn')
bullet('Tại 500 users: Error rate = 0%, P95 = 110 ms — vẫn ổn định tốt')
bullet('Tại ~900 users: Error rate = 0.01% (1/10.651 requests) — gần như hoàn hảo')
bullet('Kết luận: ✔ ĐẠT — Vượt xa ngưỡng tối thiểu 100 users', colour=GREEN)
doc.add_paragraph()

# 5.3 Error Rate / Breaking Point
heading('5.3. Breaking Point & Error Rate', level=2)
para(
    'Mục tiêu của stress test là tìm ngưỡng error rate > 5%. '
    'Sau 7 giai đoạn với tổng cộng hơn 38.000 requests, error rate tối đa ghi nhận được '
    'chỉ là 0.01% ở mức ~905 concurrent users. '
    'Server KHÔNG sập trong suốt quá trình test.'
)
bullet('Breaking point (error > 5%): KHÔNG XÁC ĐỊNH ĐƯỢC trong phạm vi test')
bullet('Mức tải tối đa đã test: ~905 concurrent users')
bullet('Đề xuất: Cần test tiếp ở mức 1.500, 2.000+ users để tìm breaking point thực sự')
bullet('Kết luận: ✔ ĐẠT (chưa tìm được điểm sập trong phạm vi test)', colour=GREEN)
doc.add_paragraph()

# 5.4 POI & Audio Loading
heading('5.4. POI Data Loading & Audio Playback (NFR: ≤ 4.000 ms)', level=2)
para(
    'Endpoint GET /api/poi và GET /api/audio/poi/{id} là hai endpoint chính trong luồng tourist. '
    'Dữ liệu tổng hợp (aggregate) cho thấy average response time luôn dưới 160 ms ở mọi giai đoạn, '
    'và P99 không vượt quá 810 ms.'
)
bullet('GET /api/poi — Avg toàn cục: ~100 ms (NFR: ≤ 4.000 ms) → ĐẠT')
bullet('GET /api/audio/poi/{id} — Avg toàn cục: ~100 ms (NFR: ≤ 4.000 ms) → ĐẠT')
bullet('Kết luận: ✔ ĐẠT', colour=GREEN)
doc.add_paragraph()

# ── 6. Quan sát đáng chú ý ──────────────────────────────────────────────────
heading('6. Quan Sát Đáng Chú Ý')

para('6.1. Hiệu năng cải thiện theo tải (Counter-intuitive)', bold=True)
para(
    'Đáng ngạc nhiên, P95 response time GIẢM khi tải tăng từ Warm-up (270 ms) xuống Cao/Rất cao (110 ms). '
    'Nguyên nhân khả năng cao: JVM warm-up (JIT compilation), Redis cache warm-up, '
    'và connection pool được tận dụng tốt hơn khi có nhiều request đồng thời.',
    indent=True
)
doc.add_paragraph()

para('6.2. RPS Plateau ở ~59 RPS', bold=True)
para(
    'Ở giai đoạn Cực đại (1) và (2) với ~900 users, RPS đạt ngưỡng ~59 req/giây và không tăng thêm '
    'dù số users tương đương. Đây là dấu hiệu server có thể đang tiếp cận giới hạn throughput nào đó '
    '(network, DB connection pool, hoặc CPU). Cần giám sát server-side metrics để xác định bottleneck.',
    indent=True
)
doc.add_paragraph()

para('6.3. P99 tăng vọt ở 905 users', bold=True)
para(
    'Tại giai đoạn Cực đại (1), P99 tăng lên 810 ms so với P95 = 200 ms — cho thấy đuôi phân phối '
    '(tail latency) bắt đầu giãn rộng, dấu hiệu của request queuing. '
    'Ở Cực đại (2), P99 giảm xuống 260 ms — có thể do connection pool đã ổn định.',
    indent=True
)
doc.add_paragraph()

# ── 7. Kết luận tổng thể ────────────────────────────────────────────────────
heading('7. Kết Luận Tổng Thể')

summary_rows = [
    [('API Response Time < 5.000 ms',   None, None), ('✔ ĐẠT',    'E2EFDA', GREEN),  ('P95 tối đa 270 ms, P99 tối đa 810 ms',              None, None)],
    [('≥ 100 Concurrent Tourist Users', None, None), ('✔ ĐẠT',    'E2EFDA', GREEN),  ('Ổn định lên đến ~905 users, error < 0.02%',          None, None)],
    [('Error Rate < 5%',                None, None), ('✔ ĐẠT',    'E2EFDA', GREEN),  ('Error rate tối đa 0.01% (1 request / 10.651)',        None, None)],
    [('POI Loading ≤ 4.000 ms',         None, None), ('✔ ĐẠT',    'E2EFDA', GREEN),  ('Avg < 160 ms ở mọi giai đoạn',                       None, None)],
    [('Audio Playback ≤ 4.000 ms',      None, None), ('✔ ĐẠT',    'E2EFDA', GREEN),  ('Avg < 160 ms ở mọi giai đoạn',                       None, None)],
    [('Breaking Point xác định',        None, None), ('⚠ CHƯA ĐỦ', 'FFF2CC', ORANGE), ('Cần test > 1.000 users để tìm điểm sập thực sự', None, None)],
]
add_table(['Tiêu chí NFR', 'Kết quả', 'Ghi chú'], summary_rows, header_fill='1F4E79')
doc.add_paragraph()

para(
    'Nhìn chung, server FlavorTales đáp ứng TẤT CẢ các Yêu Cầu Hiệu Năng được quy định trong '
    'system_design_nfr.docx §1.1 với biên độ an toàn rất lớn. Điểm cần theo dõi tiếp '
    'là bottleneck throughput khi RPS đạt ~59 req/giây ở mức ~900 users.',
    bold=False
)

# ── 8. Đề xuất tiếp theo ────────────────────────────────────────────────────
heading('8. Đề Xuất Tiếp Theo')
bullet('Chạy tiếp stress test ở mức 1.000, 1.500, 2.000 users để xác định breaking point thực sự.')
bullet('Giám sát server-side metrics (CPU, RAM, DB connections, JVM heap) song song với Locust để xác định bottleneck.')
bullet('Điều tra nguyên nhân RPS plateau ở ~59 req/giây — kiểm tra connection pool size và Nginx/reverse proxy config.')
bullet('Test thêm luồng ghi (POST like) thường xuyên hơn để đánh giá tải lên DB.')
bullet('Thiết lập alert tự động khi error rate > 1% trong môi trường production.')

# ── Save ─────────────────────────────────────────────────────────────────────
out_dir = r'd:\Codes\Seminar\Source Code\test\evaluations'
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, 'tourist-04-03-2026.docx')
doc.save(out_path)
print(f'Saved: {out_path}')
